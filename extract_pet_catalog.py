"""
Export PET swatches from DOOR Information——Gainer.xlsx (sheet "PET").

Layout: column A = color name, B = embedded picture, C = code (QZ01…), D = surface.

Run from project root (Desktop copy of the workbook is auto-detected via glob):
  python extract_pet_catalog.py

Optional: set GAINER_PET_XLSX to the full path of the workbook.
Outputs files under public/assets/catalog/pet/{Code}.png|jpg and prints TS snippet for data.ts.
"""

from __future__ import annotations

import os
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
PET_DIR = ROOT / "public" / "assets" / "catalog" / "pet"


def _slug(s) -> str:
    s = str(s).strip() if s is not None else "x"
    return re.sub(r"[^a-zA-Z0-9._-]+", "_", s)[:100] or "x"


def _resolve_workbook() -> Path | None:
    env = os.environ.get("GAINER_PET_XLSX")
    if env:
        p = Path(env).expanduser().resolve()
        if p.is_file():
            return p
        raise FileNotFoundError(f"GAINER_PET_XLSX is not a file: {p}")
    here = ROOT
    desk = Path.home() / "Desktop"
    for p in (
        here / "DOOR Information——Gainer.xlsx",
        desk / "DOOR Information——Gainer.xlsx",
        here / "DOOR Information - Gainer.xlsx",
        desk / "DOOR Information - Gainer.xlsx",
    ):
        if p.is_file():
            return p
    for folder in (here, desk):
        if folder.is_dir():
            found = sorted(folder.glob("DOOR*Gainer*.xlsx"))
            if found:
                return found[0]
    return None


def _image_at_cell(ws, excel_row: int, col_1based: int):
    for img in ws._images or []:
        fr = img.anchor._from
        if fr.row + 1 == excel_row and fr.col + 1 == col_1based:
            return img
    return None


def _best_on_row(ws, excel_row: int, prefer_cols: list[int]):
    found = []
    for img in ws._images or []:
        fr = img.anchor._from
        if fr.row + 1 != excel_row:
            continue
        found.append((fr.col + 1, img))
    if not found:
        return None
    for pc in prefer_cols:
        for col, im in found:
            if col == pc:
                return im
    found.sort(key=lambda x: x[0])
    return found[0][1]


def _save(img, code: str) -> str | None:
    try:
        ext = (img.format or "png").lower()
        if ext == "jpeg":
            ext = "jpg"
        safe = _slug(code)
        PET_DIR.mkdir(parents=True, exist_ok=True)
        out = PET_DIR / f"{safe}.{ext}"
        out.write_bytes(img._data())
        return f"/assets/catalog/pet/{safe}.{ext}"
    except (OSError, ValueError, TypeError):
        return None


def main() -> None:
    path = _resolve_workbook()
    if not path:
        raise SystemExit(
            "Could not find DOOR Information——Gainer.xlsx. "
            "Put it on Desktop or project root, or set GAINER_PET_XLSX."
        )
    wb = openpyxl.load_workbook(str(path), data_only=False)
    if "PET" not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"No sheet 'PET' in {path.name}")
    ws = wb["PET"]
    rows: list[tuple[str, str, str]] = []
    for r in range(2, ws.max_row + 1):
        name_cell = ws.cell(r, 1).value
        code_cell = ws.cell(r, 3).value
        if code_cell is None or str(code_cell).strip() in ("", "\\"):
            continue
        code = str(code_cell).strip()
        nm = str(name_cell).replace("\n", " ").strip() if name_cell else code
        im = _image_at_cell(ws, r, 2) or _best_on_row(ws, r, [2])
        if not im:
            print(f"WARN: row {r} {code} — no image in column B")
            continue
        url = _save(im, code)
        if not url:
            print(f"WARN: row {r} {code} — could not save image")
            continue
        rows.append((code, nm, url))
    wb.close()
    if not rows:
        raise SystemExit("No PET rows exported.")
    print(f"OK: {len(rows)} swatches from {path.name} -> {PET_DIR}")
    print("\nPaste into src/data.ts surfaceFinishes.pet (or re-run generate_data when 318 xlsx is present):\n")
    for code, nm, url in rows:
        esc = nm.replace("'", "\\'")
        print(
            f"    {{\n      code: '{code}',\n      name: '{esc}',\n      picture: '{url}',\n    }},"
        )


if __name__ == "__main__":
    main()
